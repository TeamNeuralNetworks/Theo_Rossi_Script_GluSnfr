import os  # Gestion des variables d'environnement et du système
# Fix pour les erreurs MKL - DOIT ÊTRE EN PREMIER
os.environ['MKL_THREADING_LAYER'] = 'GNU'  # Correction pour certains environnements numpy/scipy
os.environ['OPENBLAS_NUM_THREADS'] = '1'   # Limite le nombre de threads pour éviter les conflits
os.environ['MKL_NUM_THREADS'] = '1'

import matplotlib.pyplot as plt  # Pour les graphiques (non utilisé ici mais utile pour debug)
from scipy.optimize import curve_fit  # Pour les fits exponentiels
from scipy.signal import savgol_filter  # Pour le lissage des données
import PySimpleGUI as sg  # Interface graphique utilisateur
import numpy as np  # Calculs numériques
import pandas as pd  # Manipulation de tableaux de données
from scipy import stats  # Statistiques
import math  # Fonctions mathématiques
import random  # Tirages aléatoires
from openpyxl import Workbook  # Création de fichiers Excel
from openpyxl.utils.dataframe import dataframe_to_rows  # Conversion DataFrame -> Excel
import glob  # Recherche de fichiers
from pathlib import Path  # Gestion des chemins de fichiers

filt = 5  # Valeur de lissage (avant 9, ici 5, défini par la valeur qui donne le moins de différence pour la STD de la baseline)

# =====================
# FONCTIONS UTILITAIRES ET TRAITEMENT
# =====================

def gauss_func(mu,sigma,bins):
    '''
    Calcule une gaussienne pour un histogramme
    mu : moyenne
    sigma : écart-type
    bins : bins de l'histogramme
    '''
    y = ((1 / (np.sqrt(2 * np.pi) * sigma)) * np.exp(-0.5 * (1 / sigma * (bins - mu))**2)) 
    return y

def func_mono_exp(x, a, b, c, d):
    '''
    Fonction exponentielle décroissante pour le fit du photobleaching
    '''
    return a * np.exp(-(x-b)/c) + d

def MAD(a,axis=None):
     '''
     Calcule la médiane de la déviation absolue (robuste au bruit)
     '''
     med = np.nanmedian(a, axis=axis, keepdims=True)
     mad = np.nanmedian(np.abs(a-med),axis=axis)
     return mad

def bootstrap_patterns(patterns, run=5000, N=12, input_method='average', output_method='average'):
    '''
    Bootstrap sur des patterns synaptiques, retourne pattern médian ou moyen
    patterns : liste de tableaux (données)
    run : nombre de tirages bootstrap
    N : nombre de tirages par cycle
    input_method : 'median' ou 'average' pour chaque run
    output_method : 'median' ou 'average' sur tous les runs
    '''
    endCycle = []
    
    for i in range(run): 
        
        temp = []
        
        for j in range(N):
        
            randIndex = np.random.randint(0,len(patterns),size=1)[0]
                        
            temp.append(patterns[randIndex])
            
            if len(temp) == N:
                pass
            else:
                continue
                
        if input_method == 'median' : 
            endCycle.append(np.nanmedian(temp, axis=0))
        
        elif input_method == 'average' : 
            endCycle.append(np.nanmean(temp, axis=0))

    if output_method == 'median': 
        out_bootstrap = np.nanmedian(endCycle, axis=0) 
        out_deviation = MAD(endCycle, axis=0)
        
    elif output_method == 'average': 
        out_bootstrap = np.nanmean(endCycle, axis=0)
        out_deviation = np.nanstd(endCycle, axis=0)
        
    return np.asarray(out_bootstrap), np.asarray(out_deviation), endCycle

def load_xls(file_xls):
    '''
    Chargement sécurisé d'un fichier Excel contenant les données expérimentales.
    Tente d'ouvrir la feuille 'Traces DF_F0' avec différents moteurs, sinon la première feuille.
    Nettoie les données et extrait les colonnes temps, moyenne et sweeps.
    Retourne :
        timescale : vecteur temps
        average : trace moyenne
        SWEEPS : liste des sweeps individuelles
    '''
    try:
        print(f"Ouverture du fichier: {file_xls}")
        
        # Essayer plusieurs moteurs Excel
        engines = ['openpyxl', 'xlrd']
        df = None
        
        for engine in engines:
            try:
                df = pd.read_excel(file_xls, sheet_name='Traces DF_F0', header=0, engine=engine)
                print(f"Fichier ouvert avec le moteur: {engine}")
                break
            except Exception as e:
                print(f"Échec avec {engine}: {e}")
                continue
        
        if df is None:
            # Dernier essai avec le premier sheet
            df = pd.read_excel(file_xls, sheet_name=0, header=0)
        
        if df is None or df.empty:
            raise ValueError("Impossible d'ouvrir le fichier Excel ou fichier vide")
        
        # Nettoyer les données
        df = df.dropna(how='all')
        df = df.replace([np.inf, -np.inf], np.nan)
        
        print(f"Colonnes trouvées: {list(df.columns)}")
        print(f"Forme des données: {df.shape}")
        
        polyorder = 2 #Polyorder must be less thant filt. Before 2.
        timescale = None
        average = None
        SWEEPS = []
        
        for col_name in df.columns:
            col_data = df[col_name].dropna().values
            
            if len(col_data) < 10:
                continue
                
            if 'Time' in str(col_name) or 'time' in str(col_name).lower():
                timescale = col_data
                
            elif 'Average' in str(col_name) or 'average' in str(col_name).lower():
                average = col_data
                if len(col_data) >= filt:
                    average = savgol_filter(col_data, filt, polyorder)
                else:
                    average = col_data
                    
            else:
                if len(col_data) >= filt:
                    sweep = savgol_filter(col_data, filt, polyorder)
                else:
                    sweep = col_data
                SWEEPS.append(sweep)
        
        if timescale is None or average is None or len(SWEEPS) == 0:
            raise ValueError("Colonnes Time, Average ou sweeps manquantes")
        
        # Ajuster les longueurs
        min_len = min(len(timescale), len(average), min(len(s) for s in SWEEPS))
        timescale = timescale[:min_len]
        average = average[:min_len]
        SWEEPS = [s[:min_len] for s in SWEEPS]
        
        print(f"Données chargées: {len(timescale)} points, {len(SWEEPS)} sweeps")
        return timescale, average, SWEEPS
        
    except Exception as e:
        print(f"ERREUR lors du chargement: {e}")
        return None, None, None

def Fit_single_trace(time, trace, x_start, x_end):
    '''
    Effectue un fit exponentiel décroissant sur une fenêtre de la trace.
    Retourne :
        tau (constante de temps), paramètres du fit, indices de début/fin de la fenêtre
    '''
    
    idx_start = np.ravel(np.where(time >= x_start))[0]
    idx_stop = np.ravel(np.where(time <= x_end))[-1]
    
    x = time[idx_start:idx_stop]
    y = trace[idx_start:idx_stop]
    x2 = np.array(np.squeeze(x))
    y2 = np.array(np.squeeze(y))  
    
    try:
        param_bounds=([-np.inf,0.,0.,-1000.],[np.inf,1.,10.,1000.])      # bornes pour le fit, ok for seconds. If milliseconds, change param 2 and 3.
        popt, pcov = curve_fit(func_mono_exp, x2, y2, bounds=param_bounds, maxfev=10000) 
        return popt[2], popt, idx_start, idx_stop
    except:
        print ('Fit failed')
        popt = [float('nan')] * 4
        return float('nan'), popt, idx_start, idx_stop

def bleaching_correction(time, trace, start, stop):
    '''
    Correction du photobleaching par soustraction d'un fit exponentiel sur la fenêtre [start, stop].
    Fonctionne sur une trace ou une liste de traces.
    '''
    
    x1 = float(start)
    x2 = float(stop)
    
    if len(trace) == 0:
        print('Bleaching failed: empty window')
        
    elif type(trace) == np.ndarray:
        tau, popt, xstart, xstop = Fit_single_trace(time, trace, x1, x2)
        bleach = func_mono_exp(time, *popt)
        no_bleach = trace-bleach

    elif type(trace) == list:
        param = [Fit_single_trace(time, trace[i], x1, x2) for i in range(len(trace))]
        bleach = [func_mono_exp(time, *param[i][1]) for i in range(len(trace))]
        no_bleach = [trace[i]-bleach[i] for i in range(len(trace))]
        print('bleaching correction is done')
    return no_bleach

def leak_subtraction(time, trace, start, stop):
    ''' 
    Soustraction du leak (offset) sur la fenêtre [start, stop].
    Fonctionne sur une trace ou une liste de traces.
    '''
    
    x1 = np.ravel(np.where(time >= float(start)))[0]
    x2 = np.ravel(np.where(time <= float(stop)))[-1]
    
    if len(trace) == 0:
        print('Leak failed: empty window')
        
    elif type(trace) == np.ndarray:
        leak = np.mean(trace[x1:x2])
        no_leak = trace-leak
        
    elif type(trace) == list:
        leak = [np.mean(trace[i][x1:x2]) for i in range(len(trace))]
        no_leak = [trace[i]-leak[i] for i in range(len(trace))]
        
        print('substraction is done')
    return no_leak

def residual_sublimation(time, trace, start, stop, freq, n_peaks):
    ''' 
    Soustraction du résiduel précédant chaque pic.
    freq : fréquence de stimulation (str)
    n_peaks : nombre de pics à traiter
    Fonctionne sur une trace ou une liste de traces.
    '''
    
    if len(trace) == 0:
        print('Res sublimation failed: empty window')
        
    elif type(trace) == np.ndarray:
        x1 = np.ravel(np.where(time >= start))[0]
        x2 = np.ravel(np.where(time <= stop))[-1]
        x2_bis = np.ravel(np.where(time <= stop))[-1]
        bsl = trace[0:x2]
        PEAKS_NO_RES = []
        for i in range(n_peaks):
            res = np.mean(trace[x1:x2])
            if freq == '20':
                stop += 0.05
                x2_bis = np.ravel(np.where(time <= stop))[-1]
                peak_no_res = trace[x2:x2_bis]-res
                start += 0.05
            elif freq == '50':  
                stop += 0.02
                x2_bis = np.ravel(np.where(time <= stop))[-1]
                peak_no_res = trace[x2:x2_bis]-res
                start += 0.02
            elif freq == '100':  
                stop += 0.01
                x2_bis = np.ravel(np.where(time <= stop))[-1]
                peak_no_res = trace[x2:x2_bis]-res
                start += 0.01
                
            x1 = np.ravel(np.where(time >= start))[0]
            x2 = np.ravel(np.where(time <= stop))[-1]
            PEAKS_NO_RES.append(peak_no_res)
            
        end = np.append(np.hstack(PEAKS_NO_RES), trace[x2:])
        no_res = np.append(bsl, end)
        
    elif type(trace) == list:
        x2_bsl = np.ravel(np.where(time <= stop))[-1]
        no_res = []
        for i in range(len(trace)):
            alpha = start
            omega = stop
            x1 = np.ravel(np.where(time >= alpha))[0]
            x2 = np.ravel(np.where(time <= omega))[-1]
            x2_bis = np.ravel(np.where(time <= omega))[-1]
            bsl = trace[i][0:x2_bsl]
            PEAKS_NO_RES = []
            for j in range(n_peaks):
                res = np.mean(trace[i][x1:x2])
                if freq == '20':
                    omega += 0.05
                    x2_bis = np.ravel(np.where(time <= omega))[-1]
                    peak_no_res = trace[i][x2:x2_bis]-res
                    alpha += 0.05
                elif freq == '50':  
                    omega += 0.02
                    x2_bis = np.ravel(np.where(time <= omega))[-1]
                    peak_no_res = trace[i][x2:x2_bis]-res
                    alpha += 0.02
                elif freq == '100':
                    omega += 0.01
                    x2_bis = np.ravel(np.where(time <= omega))[-1]
                    peak_no_res = trace[i][x2:x2_bis]-res
                    alpha += 0.01
                
                x1 = np.ravel(np.where(time >= alpha))[0]
                x2 = np.ravel(np.where(time <= omega))[-1]
                PEAKS_NO_RES.append(peak_no_res)
 
            end = np.append(np.hstack(PEAKS_NO_RES), trace[i][x2:])
            no_res_trace = np.append(bsl, end)
            no_res.append(no_res_trace)
        print('sublimation is done')
    return no_res

def values_extraction(time, trace, start, stop):    
    '''
    Extrait les valeurs d'une fenêtre temporelle [start, stop] d'une trace ou liste de traces.
    '''
    
    x1 = np.ravel(np.where(time >= float(start)))[0]
    x2 = np.ravel(np.where(time <= float(stop)))[-1]
    
    if len(trace) == 0:
        print('Extraction failed: empty window')
        
    elif type(trace) == np.ndarray:
        windows = trace[x1:x2]
        
    elif type(trace) == list:
        windows = [trace[i][x1:x2] for i in range(len(trace))]
        
    return windows

def process_single_file(file_path, output_folder, parameters):
    """
    Traite un seul fichier avec les paramètres donnés
    Applique les corrections (photobleaching, leak, résiduel), extrait les fenêtres d'intérêt,
    effectue le bootstrap et sauvegarde les résultats dans des fichiers Excel.
    Retourne True, nom du fichier, liste des pourcentages d'échec (PercFail) pour chaque pic.
    """
    try:
        print(f"\n=== Traitement du fichier: {file_path} ===")
        # Chargement du fichier et extraction des données principales
        time, avg, sweeps = load_xls(file_path)
        if time is None or avg is None or sweeps is None:
            print(f"Erreur lors du chargement de {file_path}")
            return False
        
        # Application des corrections automatiques
        print("Application des corrections...")
        
        # 1. Correction du photobleaching (décroissance exponentielle)
        avg_corrected = bleaching_correction(time, avg, parameters['photo_start'], parameters['photo_stop'])
        sweeps_corrected = bleaching_correction(time, sweeps, parameters['photo_start'], parameters['photo_stop'])
        
        # 2. Soustraction du leak (offset)
        avg_corrected = leak_subtraction(time, avg_corrected, parameters['leak_start'], parameters['leak_stop'])
        sweeps_corrected = leak_subtraction(time, sweeps_corrected, parameters['leak_start'], parameters['leak_stop'])
        
        # 3. Sublimation résiduelle (soustraction du résiduel avant chaque pic)
        avg_corrected = residual_sublimation(time, avg_corrected, parameters['res_start'], parameters['res_stop'], 
                                           parameters['frequency'], parameters['n_peaks'])
        sweeps_corrected = residual_sublimation(time, sweeps_corrected, parameters['res_start'], parameters['res_stop'], 
                                              parameters['frequency'], parameters['n_peaks'])
        
        # 4. Extraction des fenêtres temporelles pour l'analyse des pics et du bruit
        file_name = Path(file_path).stem
        a = parameters['peak_start']
        b = parameters['peak_stop']
        WINDOWS_NS, WINDOWS_AMP = [], []
        for i in range(parameters['n_peaks']):
            windows_amp = values_extraction(time, sweeps_corrected, a, b)
            windows_ns = values_extraction(time, sweeps_corrected, parameters['noise_start'], parameters['noise_stop'])
            # Décalage de la fenêtre selon la fréquence
            if parameters['frequency'] == '20':
                a += 0.05
                b += 0.05
            elif parameters['frequency'] == '50':  
                a += 0.02
                b += 0.02
            elif parameters['frequency'] == '100':  
                a += 0.01
                b += 0.01
            WINDOWS_NS.append(windows_ns)
            WINDOWS_AMP.append(windows_amp)
        
        # 5. Analyse bootstrap sur chaque pic
        print("Analyse bootstrap...")
        wb_hist = Workbook()  # Fichier Excel pour les cycles bootstrap
        wb_data = Workbook()  # Fichier Excel pour les résultats agrégés
        ALL_FAILS = []  # Pourcentage d'échec (zscore <= 2) pour chaque pic
        for peak in range(parameters['n_peaks']):
            EP, PEAK, ZSCORE, AMP, NS_AMP, STD_NS, STD_AMP = [], [], [], [], [], [], []
            df_episodes = pd.DataFrame(index=None, columns=None)
            print(f'BOOTSTRAP PEAK{peak+1}')
            for item in range(len(WINDOWS_AMP[peak])):
                EP.append(f'Episode {item}')
                # Bootstrap sur la fenêtre du pic
                amp_mean, amp_std, amp_cycle = bootstrap_patterns(WINDOWS_AMP[peak][item], N=len(WINDOWS_AMP[peak][item]))
                # Bootstrap bruit (tirages aléatoires dans la fenêtre bruit)
                ns_cycle = [np.mean(random.sample(windows_ns[item].tolist(), len(WINDOWS_AMP[peak][item]))) 
                           for i in range(len(amp_cycle))]
                ns_mean = np.mean(ns_cycle)
                ns_std = np.std(ns_cycle)
                NS_AMP.append(ns_mean)
                STD_AMP.append(amp_std)
                STD_NS.append(ns_std)
                amp = amp_mean - ns_mean
                AMP.append(amp)
                df_episodes[f'noise{item}'] = ns_cycle
                df_episodes[f'amp{item}'] = amp_cycle
                z_score = amp/ns_std
                ZSCORE.append(z_score)
                if z_score <= 2:
                    PEAK.append('Fail')
                else:
                    PEAK.append('Success')
            # Calcul du pourcentage d'échec pour ce pic
            percentage_failures = (PEAK.count('Fail')/len(WINDOWS_AMP[peak]))*100
            ALL_FAILS.append(percentage_failures)
            # Tableaux de résultats pour Excel
            df_ep = pd.concat((pd.DataFrame(EP), pd.DataFrame(PEAK), pd.DataFrame(AMP), pd.DataFrame(NS_AMP),
                               pd.DataFrame(STD_AMP), pd.DataFrame(STD_NS), pd.DataFrame(ZSCORE)), axis=1)
            df_ep.columns = ['Episode','PEAK','AMP','AMPns','Std_amp','Std_ns','Zscore']
            df_ep['PercFail'] = percentage_failures
            sheet_hist = wb_hist.create_sheet(f'PEAK{peak+1}')
            sheet_data = wb_data.create_sheet(f'PEAK{peak+1}')
            for r in dataframe_to_rows(df_episodes, index=False, header=True):
                sheet_hist.append(r)
            for r in dataframe_to_rows(df_ep, index=False, header=True):
                sheet_data.append(r)
        
        # 6. Sauvegarde des fichiers Excel individuels
        wb_hist.remove(wb_hist['Sheet'])
        wb_data.remove(wb_data['Sheet'])
        hist_path = os.path.join(output_folder, f'{file_name}_histograms_bootstrap.xlsx')
        data_path = os.path.join(output_folder, f'{file_name}_data_bootstrap.xlsx')
        wb_hist.save(hist_path)
        wb_data.save(data_path)
        print(f"Fichier traité avec succès: {file_name}")
        print(f"Fichiers sauvegardés: {hist_path}, {data_path}")
        
        # 7. Retourne True, nom du fichier, et liste des pourcentages d'échec pour le récapitulatif
        return True, file_name, ALL_FAILS
    except Exception as e:
        print(f"Erreur lors du traitement de {file_path}: {e}")
        return False, None, None

def batch_process():
    """
    Interface graphique pour le traitement par lots de fichiers Excel.
    Permet de sélectionner les dossiers, de paramétrer les corrections et l'analyse,
    lance le traitement de tous les fichiers trouvés, affiche la progression et sauvegarde un récapitulatif.
    """
    sg.theme('DarkBlue')
    
    # Définition de la fenêtre principale avec tous les paramètres utilisateur
    layout = [
      [sg.Text('Dossier contenant les fichiers Excel:')],
      [sg.InputText(size=(50,1), key='input_folder'), sg.FolderBrowse()],
      [sg.Text('Dossier de sortie:')],
      [sg.InputText(size=(50,1), key='output_folder'), sg.FolderBrowse()],
      [sg.Frame('Paramètres de correction', [
          [sg.Text('Photobleaching:'), sg.InputText('0.01', size=(6,1), key='photo_start'), 
           sg.Text('à'), sg.InputText('0.95', size=(6,1), key='photo_stop'), sg.Text('sec')],
          [sg.Text('Leak:'), sg.InputText('0.85', size=(6,1), key='leak_start'), 
           sg.Text('à'), sg.InputText('0.95', size=(6,1), key='leak_stop'), sg.Text('sec')],
          [sg.Text('Résiduel:'), sg.InputText('0.99', size=(6,1), key='res_start'), 
           sg.Text('à'), sg.InputText('1.0', size=(6,1), key='res_stop'), sg.Text('sec')]
      ])],
      [sg.Frame('Paramètres d\'analyse', [
          [sg.Text('Fréquence (Hz):'), sg.InputText('20', size=(4,1), key='frequency'),
           sg.Text('Nombre de pics:'), sg.InputText('3', size=(4,1), key='n_peaks')],
          [sg.Text('Fenêtre pic:'), sg.InputText('0.998', size=(6,1), key='peak_start'), 
           sg.Text('à'), sg.InputText('1.020', size=(6,1), key='peak_stop')],
          [sg.Text('Fenêtre bruit:'), sg.InputText('0.6', size=(6,1), key='noise_start'), 
           sg.Text('à'), sg.InputText('0.9', size=(6,1), key='noise_stop')]
      ])],
      [sg.Button('Traiter tous les fichiers', size=(20,2)), sg.Button('Quitter')]
    ]


    window = sg.Window('Bootstrap Batch Processing', layout)


    while True:
        event, values = window.read()

        if event in (None, 'Quitter'):
            break

        if event == 'Traiter tous les fichiers':
            input_folder = values['input_folder']
            output_folder = values['output_folder']
            
            # Vérification des dossiers
            if not input_folder or not output_folder:
                sg.popup_error('Veuillez sélectionner les dossiers d\'entrée et de sortie')
                continue
            
            if not os.path.exists(output_folder):
                os.makedirs(output_folder)
            
            # Recherche des fichiers Excel dans le dossier d'entrée
            excel_files = []
            for ext in ['*.xlsx', '*.xls']:
                excel_files.extend(glob.glob(os.path.join(input_folder, ext)))
            
            if not excel_files:
                sg.popup_error('Aucun fichier Excel trouvé dans le dossier d\'entrée')
                continue
            
            # Récupération des paramètres utilisateur
            parameters = {
                'photo_start': float(values['photo_start']),
                'photo_stop': float(values['photo_stop']),
                'leak_start': float(values['leak_start']), 
                'leak_stop': float(values['leak_stop']),
                'res_start': float(values['res_start']),
                'res_stop': float(values['res_stop']),
                'frequency': values['frequency'],
                'n_peaks': int(values['n_peaks']),
                'peak_start': float(values['peak_start']),
                'peak_stop': float(values['peak_stop']),
                'noise_start': float(values['noise_start']),
                'noise_stop': float(values['noise_stop'])
            }
            
            # Traitement des fichiers
            total_files = len(excel_files)
            successful = 0
            failed = 0
            all_in_results = []  # Pour stocker les résultats pour All_In.xlsx
            
            # Fenêtre de progression
            progress_layout = [[sg.Text(f'Traitement en cours... 0/{total_files}', key='progress_text')],
                              [sg.ProgressBar(total_files, orientation='h', size=(50, 20), key='progress_bar')],
                              [sg.Button('Annuler', key='cancel')]]
            progress_window = sg.Window('Progression', progress_layout, finalize=True)
            
            # Boucle de traitement de chaque fichier Excel
            for i, file_path in enumerate(excel_files):
                event_prog, values_prog = progress_window.read(timeout=10)
                if event_prog == 'cancel':
                    break
                
                progress_window['progress_text'].update(f'Traitement: {os.path.basename(file_path)} ({i+1}/{total_files})')
                progress_window['progress_bar'].update(i)
                
                # Appel du traitement sur un fichier
                result, file_name, all_fails = process_single_file(file_path, output_folder, parameters)
                if result:
                    successful += 1
                    # On ne garde que les 3 premiers PercFail (PEAK1, PEAK2, PEAK3)
                    row = [file_name]
                    for idx in range(3):
                        if all_fails and len(all_fails) > idx:
                            row.append(all_fails[idx])
                        else:
                            row.append('')
                    all_in_results.append(row)
                else:
                    failed += 1
            
            progress_window.close()
            
            # Sauvegarde du fichier All_In.xlsx récapitulatif
            if all_in_results:
                df_allin = pd.DataFrame(all_in_results, columns=['Fichier', 'PercFail_PEAK1', 'PercFail_PEAK2', 'PercFail_PEAK3'])
                allin_path = os.path.join(output_folder, 'All_In.xlsx')
                df_allin.to_excel(allin_path, index=False)
            
            # Affichage du résumé final
            sg.popup(f'Traitement terminé!\n\n'
                    f'Fichiers traités avec succès: {successful}\n'
                    f'Fichiers échoués: {failed}\n'
                    f'Total: {total_files}')

    window.close()

if __name__ == '__main__':
    batch_process()