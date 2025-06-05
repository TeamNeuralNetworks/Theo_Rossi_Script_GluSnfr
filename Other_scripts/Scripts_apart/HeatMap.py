# -*- coding: utf-8 -*-
"""
Created on Fri Apr 16 17:43:18 2021

@author: Theo.ROSSI
"""


import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy.stats.kde import gaussian_kde
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows


path = r'E:\AAVDJ.GluSnFR-S72A\AAVDJ.GluSnFR-S72A_Good_CaMg_ratio\GluSnFR_avg_amps2.xlsx'
df = pd.read_excel(path)
individuals = df['Individuals']
targets = df['Targets']
amps = df.iloc[:,1:11]
ppr = pd.DataFrame([amps.iloc[:,i]/amps.iloc[:,0] for i in range(amps.shape[1])]).T

arr_x = np.ravel([[i+1 for j in range(amps.shape[0])] for i in range(amps.shape[1])])
arr_y = np.ravel([amps.iloc[:,i].values.tolist() for i in range(amps.shape[1])])
# arr_x, arr_y = np.meshgrid(np.linspace(1,10,10), np.linspace(0,2,21), sparse=False, indexing='ij')

Z, xedges, yedges = np.histogram2d(arr_x, arr_y, bins=(10,40), range=[[1,10],[0,2]])
k = gaussian_kde(np.vstack([arr_x, arr_y]))
xi, yi = np.mgrid[arr_x.min():arr_x.max():arr_x.size**0.5*1j,arr_y.min():arr_y.max():arr_y.size**0.5*1j]
zi = k(np.vstack([xi.flatten(), yi.flatten()]))

IN = pd.DataFrame([amps.iloc[i,:].T for i in range(targets.shape[0]) if targets[i] == 0.0])
IN_ppr = pd.DataFrame([ppr.iloc[i,:].T for i in range(targets.shape[0]) if targets[i] == 0.0])
PC = pd.DataFrame([amps.iloc[i,:].T for i in range(targets.shape[0]) if targets[i] == 1.0])
PC_ppr = pd.DataFrame([ppr.iloc[i,:].T for i in range(targets.shape[0]) if targets[i] == 1.0])
Unknown = pd.DataFrame([amps.iloc[i,:].T for i in range(targets.shape[0]) if targets[i] != 0.0 and targets[i] != 1.0])
Unknown_ppr = pd.DataFrame([ppr.iloc[i,:].T for i in range(targets.shape[0]) if targets[i] != 0.0 and targets[i] != 1.0])

Labels = ['Interneurons', 'Purkinje', 'Unknown']
Groups_amps = [IN, PC, Unknown]
Groups_ppr = [IN_ppr, PC_ppr, Unknown_ppr]




# wb_amps = Workbook()
# wb_ppr = Workbook()

# for i in range(len(Labels)):
#     Groups_amps[i].columns = ['AMP1','AMP2','AMP3','AMP4','AMP5','AMP6','AMP7','AMP8','AMP9','AMP10']
#     Groups_amps[i]['Individuals'] = individuals
#     Groups_ppr[i].columns = ['PPR1/1','PPR2/1','PPR3/1','PPR4/1','PPR5/1','PPR6/1','PPR7/1','PPR8/1','PPR9/1','PPR10/1']
#     Groups_ppr[i]['Individuals'] = individuals
    
#     sheet_amps = wb_amps.create_sheet('{}'.format(Labels[i]))
#     sheet_ppr = wb_ppr.create_sheet('{}'.format(Labels[i]))
    
#     for r in dataframe_to_rows(Groups_amps[i], index=False, header=True):
#         sheet_amps.append(r)
#     for r in dataframe_to_rows(Groups_ppr[i], index=False, header=True):
#         sheet_ppr.append(r)

# wb_amps.remove(wb_amps['Sheet'])
# wb_ppr.remove(wb_ppr['Sheet'])

# wb_amps.save('E:\AAVDJ.GluSnFR-S72A\AAVDJ.GluSnFR-S72A_Good_CaMg_ratio\iGluSnFR_amps_targets.xlsx')
# wb_ppr.save('E:\AAVDJ.GluSnFR-S72A\AAVDJ.GluSnFR-S72A_Good_CaMg_ratio\iGluSnFR_ppr_targets.xlsx')





fig1, ax1 = plt.subplots(1,3, tight_layout=True)
ax1[0].pcolormesh(xedges, yedges, Z.T, cmap='magma')
ax1[1].pcolormesh(xi, yi, zi.reshape(xi.shape), cmap='magma')
ax1[2].contourf(xi,yi,zi.reshape(xi.shape), cmap='magma')


fig2, ax2 = plt.subplots(1, 2, figsize=(12,6), tight_layout=True)

fig3 = plt.figure(figsize=(12,6), tight_layout=True)
gs3 = fig3.add_gridspec(2,3)
ax3_1 = fig3.add_subplot(gs3[0,0])
ax3_2 = fig3.add_subplot(gs3[0,1], sharey=ax3_1)
ax3_3 = fig3.add_subplot(gs3[0,2], sharey=ax3_1)
ax3_1_1 = fig3.add_subplot(gs3[1,0])
ax3_2_2 = fig3.add_subplot(gs3[1,1], sharey=ax3_1_1)
ax3_3_3 = fig3.add_subplot(gs3[1,2], sharey=ax3_1_1)

ax2[0].plot(IN.iloc[:,:10].T, color='b', linewidth=8, alpha=0.1)
ax2[1].plot(IN_ppr.iloc[:,:10].T, color='b', linewidth=8, alpha=0.1)
ax3_1.plot(IN.iloc[:,:10].T, color='b', linewidth=8, alpha=0.1)
ax3_1.plot(np.mean(IN.iloc[:,:10], axis=0), color='b', linewidth=2)
ax3_1_1.plot(IN_ppr.iloc[:,:10].T, color='b', linewidth=8, alpha=0.1)
ax3_1_1.plot(np.mean(IN_ppr.iloc[:,:10], axis=0), color='b', linewidth=2)

ax2[0].plot(PC.iloc[:,:10].T, color='r', linewidth=8, alpha=0.1)
ax2[1].plot(PC_ppr.iloc[:,:10].T, color='r', linewidth=8, alpha=0.1)
ax3_2.plot(PC.iloc[:,:10].T, color='r', linewidth=8, alpha=0.1)
ax3_2.plot(np.mean(PC.iloc[:,:10], axis=0), color='r', linewidth=2)
ax3_2_2.plot(PC_ppr.iloc[:,:10].T, color='r', linewidth=8, alpha=0.1)
ax3_2_2.plot(np.mean(PC_ppr.iloc[:,:10], axis=0), color='r', linewidth=2)

ax2[0].plot(Unknown.iloc[:,:10].T, color='k', linewidth=8, alpha=0.1)
ax2[1].plot(Unknown_ppr.iloc[:,:10].T, color='k', linewidth=8, alpha=0.1)
ax3_3.plot(Unknown.iloc[:,:10].T, color='k', linewidth=8, alpha=0.1)
ax3_3.plot(np.mean(Unknown.iloc[:,:10], axis=0), color='k', linewidth=2)
ax3_3_3.plot(Unknown_ppr.iloc[:,:10].T, color='k', linewidth=8, alpha=0.1)
ax3_3_3.plot(np.mean(Unknown_ppr.iloc[:,:10], axis=0), color='k', linewidth=2)

leg = ax2[0].legend(Labels)
leg.legendHandles[0].set_color('b')
leg.legendHandles[1].set_color('r')
leg.legendHandles[2].set_color('k')
